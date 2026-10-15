"""Read explicit source manifests; keep invalid source rows and immutable snapshots."""

import csv
import json
from pathlib import Path

import numpy as np
import pandas as pd

from src.geometry.triangle import LENGTH_TO_UM, feature_frame, inside, vertices
from utils.provenance import file_sha256, write_json

from .contracts import ORIGINAL_ORIGINS, SCHEMA_VERSION, TargetSpec


def read_source(path):
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            headers = list(next(book.active.iter_rows(values_only=True)))
        finally:
            book.close()
    elif path.suffix.lower() in {".csv", ".tsv"}:
        separator = "\t" if path.suffix.lower() == ".tsv" else ","
        with path.open(encoding="utf-8-sig", newline="") as stream:
            headers = next(csv.reader(stream, delimiter=separator))
    else:
        raise ValueError("Supported source formats: xlsx, csv, tsv")
    if not headers or any(h is None or not str(h).strip() for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("Headers must be nonempty and distinct before parsing")
    return pd.read_excel(path) if path.suffix.lower() == ".xlsx" else pd.read_csv(path, sep=separator)


def audit_sources(manifest_path, out_dir):
    manifest_path, out_dir = Path(manifest_path), Path(out_dir)
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("schema_version") != SCHEMA_VERSION or not manifest.get("sources"):
        raise ValueError("Expected schema_version=1 and nonempty sources")
    if manifest.get("classes", ["UF", "F"]) not in (["UF", "F"], ["UF", "F", "WF"]):
        raise ValueError("classes must be [UF,F] or [UF,F,WF]")
    frames, rejected, issues, source_hashes, targets = [], [], [], {}, {}
    for number, spec in enumerate(manifest["sources"]):
        path = (manifest_path.parent / spec["path"]).resolve()
        try:
            task = spec["task"]
            if task not in {"FD", "FL", "Fzong"}:
                raise ValueError("task must be FD, FL or Fzong")
            target = TargetSpec(**spec["target"])
            if spec["length_unit"] not in LENGTH_TO_UM or spec["radius_unit"] not in LENGTH_TO_UM:
                raise ValueError("Explicit length_unit and radius_unit required")
            if spec.get("coordinate_frame") != "global_xz":
                raise ValueError("Sources must declare global_xz; no implicit coordinate transforms")
            if spec.get("value_origin") not in ORIGINAL_ORIGINS:
                raise ValueError("Only independent measured/simulated labels accepted in baseline snapshots")
            triangle = vertices(spec["geometry"])
            before = file_sha256(path)
            raw = read_source(path)
            if file_sha256(path) != before:
                raise ValueError("Source changed during read")
            source_hashes[str(path)] = before
            if raw.empty or raw.columns.duplicated().any():
                raise ValueError("Source must be nonempty with distinct headers")
            mapping = spec.get("columns", {})
            names = [mapping.get(c, c) for c in ("x", "z", "re")]
            if len(set([*names, target.column])) != 4 or (
                spec.get("radius") is None and mapping.get("radius", "R") in [*names, target.column]
            ):
                raise ValueError("Coordinate/Re/target columns must be distinct")
            values = raw[names + [target.column]].apply(pd.to_numeric, errors="coerce")
            values.columns = ["x", "z", "re", "target"]
            values[["x", "z"]] *= LENGTH_TO_UM[spec["length_unit"]]
            radius = spec.get("radius")
            values["radius_um"] = (
                pd.to_numeric(raw[mapping.get("radius", "R")], errors="coerce") if radius is None else float(radius)
            ) * LENGTH_TO_UM[spec["radius_unit"]]
            valid = np.isfinite(values.to_numpy()).all(axis=1) & values.re.gt(0) & values.radius_um.gt(0)
            interior = np.zeros(len(values), dtype=bool)
            interior[valid] = inside(values.loc[valid, ["x", "z"]].to_numpy(), triangle)
            valid &= interior
            bad = raw.loc[~valid].copy()
            if len(bad):
                bad["source_file"], bad["source_row"], bad["reason"] = (
                    str(path),
                    bad.index + 2,
                    "invalid_numeric_or_outside_geometry",
                )
                rejected.append(bad)
                issues.append({"source": str(path), "severity": "REJECT", "reason": "invalid_rows", "count": len(bad)})
            target_dict = target.to_dict()
            # Source column names may differ; the physical target contract may not.
            semantic = {k: v for k, v in target_dict.items() if k != "column"}
            if task in targets and targets[task] != semantic:
                raise ValueError("Conflicting target definitions for " + task)
            targets[task] = semantic
            source_frames = []
            for (re_value, radius_um), part in values.loc[valid].groupby(["re", "radius_um"], sort=True):
                frame = feature_frame(part[["x", "z"]].to_numpy(), triangle, re_value, radius_um, spec["experiment_id"])
                for column in ("ab", "ac", "bc", "h", "area", "perimeter", "angle_a", "angle_b", "angle_c"):
                    original = mapping.get(column, column)
                    if original not in raw:
                        continue
                    supplied = pd.to_numeric(raw.loc[part.index, original], errors="coerce").to_numpy()
                    factor = (
                        1.0
                        if column.startswith("angle_")
                        else LENGTH_TO_UM[spec["length_unit"]] ** (2 if column == "area" else 1)
                    )
                    if not np.allclose(supplied * factor, frame[column].to_numpy(), rtol=1e-4, atol=1e-6):
                        raise ValueError("Supplied geometry conflicts with declared vertices: " + column)
                frame["target"] = part.target.to_numpy()
                frame["task"], frame["source_file"], frame["source_hash"] = task, str(path), before
                frame["source_row"], frame["value_origin"] = part.index.to_numpy() + 2, spec["value_origin"]
                source_frames.append(frame)
            frames.extend(source_frames)
            if target.quantity == "unverified":
                issues.append(
                    {
                        "source": str(path),
                        "severity": "REVIEW",
                        "reason": "unverified_target_physics",
                        "count": len(values),
                    }
                )
        except (KeyError, ValueError, OSError, TypeError, StopIteration) as error:
            issues.append({"source": str(path), "severity": "REJECT", "reason": str(error), "count": 1})
    if not frames:
        points = pd.DataFrame(
            columns=["sample_id", "point_id", "task", "geometry_family_id", "experiment_id", "re", "radius_um"]
        )
    else:
        points = pd.concat(frames, ignore_index=True)
        duplicates = points.duplicated(["sample_id", "point_id", "task"], keep=False)
        if duplicates.any():
            rejected.append(points.loc[duplicates].assign(reason="duplicate_point_identity"))
            issues.append(
                {
                    "source": "dataset",
                    "severity": "REJECT",
                    "reason": "duplicate_point_identity",
                    "count": int(duplicates.sum()),
                }
            )
            points = points.loc[~duplicates].reset_index(drop=True)
        ambiguous = points.groupby(["experiment_id", "re", "radius_um"]).sample_id.nunique().gt(1)
        if ambiguous.any():
            issues.append(
                {
                    "source": "dataset",
                    "severity": "REJECT",
                    "reason": "experiment_id_maps_to_multiple_geometries",
                    "count": int(ambiguous.sum()),
                }
            )
    samples = points[["sample_id", "experiment_id", "geometry_family_id", "re", "radius_um"]].drop_duplicates()
    labels = manifest.get("labels", [])
    if labels:
        label_frame = pd.DataFrame(labels)
        key = ["experiment_id", "re", "radius_um"]
        try:
            if not manifest.get("label_provenance"):
                raise ValueError("Independent label_provenance is required")
            classes = manifest.get("classes", ["UF", "F"])
            if classes not in (["UF", "F"], ["UF", "F", "WF"]):
                raise ValueError("classes must be [UF,F] or [UF,F,WF]")
            if label_frame.duplicated(key).any() or not label_frame.label.isin(classes).all():
                raise ValueError("Duplicate or unsupported labels")
            for column in ("re", "radius_um"):
                label_frame[column] = pd.to_numeric(label_frame[column], errors="raise")
                if not np.isfinite(label_frame[column]).all() or label_frame[column].le(0).any():
                    raise ValueError("Label Re and radius_um must be positive finite values")
            joined = samples.merge(
                label_frame[key + ["label"]], on=key, how="outer", indicator=True, validate="one_to_one"
            )
            if joined._merge.eq("right_only").any():
                raise ValueError("Labels without an audited sample")
            samples = joined.drop(columns="_merge")
        except (ValueError, KeyError) as error:
            issues.append({"source": "labels", "severity": "REJECT", "reason": str(error), "count": len(labels)})
    out_dir.mkdir(parents=True, exist_ok=True)
    points.to_csv(out_dir / "points.csv", index=False)
    samples.to_csv(out_dir / "samples.csv", index=False)
    pd.concat(rejected, ignore_index=True).to_csv(
        out_dir / "rejected_rows.csv", index=False
    ) if rejected else pd.DataFrame(columns=["reason"]).to_csv(out_dir / "rejected_rows.csv", index=False)
    pd.DataFrame(issues, columns=["source", "severity", "reason", "count"]).to_csv(out_dir / "issues.csv", index=False)
    status = (
        "REJECT"
        if not len(points) or any(i["severity"] == "REJECT" for i in issues)
        else "REVIEW"
        if issues
        else "ACCEPT"
    )
    write_json(
        out_dir / "audit.json",
        {"status": status, "accepted_rows": len(points), "samples": len(samples), "issues": issues},
    )
    snapshot = {
        "schema_version": SCHEMA_VERSION,
        "status": status,
        "targets": targets,
        "classes": manifest.get("classes", ["UF", "F"]),
        "source_hashes": source_hashes,
        "manifest": manifest,
        "coordinate_frame": "global_xz",
        "length_unit": "um",
        "files": {name: file_sha256(out_dir / name) for name in ("points.csv", "samples.csv")},
    }
    write_json(out_dir / "snapshot.json", snapshot)
    return snapshot


def load_snapshot(path):
    path = Path(path)
    meta = json.loads((path / "snapshot.json").read_text(encoding="utf-8"))
    if meta.get("schema_version") != SCHEMA_VERSION or meta["status"] == "REJECT":
        raise ValueError("Snapshot rejected or unsupported; inspect audit.json")
    if set(meta["files"]) != {"points.csv", "samples.csv"}:
        raise ValueError("Unexpected snapshot files")
    for name, digest in meta["files"].items():
        if file_sha256(path / name) != digest:
            raise ValueError("Snapshot content hash mismatch: " + name)
    return pd.read_csv(path / "points.csv"), pd.read_csv(path / "samples.csv"), meta
